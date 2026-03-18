"""Parse documents from raw bytes and split them into LangChain chunks."""

from __future__ import annotations

import io
import logging
from pathlib import Path

from langchain_core.documents import Document
from langchain_text_splitters.character import RecursiveCharacterTextSplitter

from config import Config

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Converts raw file bytes into a list of LangChain ``Document`` objects."""

    def __init__(self) -> None:
        self._splitter = RecursiveCharacterTextSplitter(
            chunk_size=Config.CHUNK_SIZE,
            chunk_overlap=Config.CHUNK_OVERLAP,
            separators=["\n\n", "\n", ". ", " ", ""],
        )

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def process(self, file_bytes: bytes, filename: str) -> list[Document]:
        """
        Parse *file_bytes* according to the file extension of *filename* and
        return a list of chunked ``Document`` objects ready for embedding.
        """
        ext = Path(filename).suffix.lower()
        text = self._extract_text(file_bytes, ext, filename)
        if not text.strip():
            raise ValueError(f"No text could be extracted from '{filename}'.")

        chunks = self._splitter.create_documents(
            texts=[text],
            metadatas=[{"source": filename}],
        )
        logger.info("'%s' → %d chunks", filename, len(chunks))
        return chunks

    # ------------------------------------------------------------------
    # Private parsers
    # ------------------------------------------------------------------

    def _extract_text(self, file_bytes: bytes, ext: str, filename: str) -> str:
        parsers = {
            ".pdf": self._parse_pdf,
            ".txt": self._parse_txt,
            ".md": self._parse_txt,
            ".docx": self._parse_docx,
            ".csv": self._parse_csv,
            ".xlsx": self._parse_xlsx,
        }
        parser = parsers.get(ext)
        if parser is None:
            raise ValueError(f"No parser available for extension '{ext}'.")
        return parser(file_bytes)

    @staticmethod
    def _parse_pdf(file_bytes: bytes) -> str:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n\n".join(pages)

    @staticmethod
    def _parse_txt(file_bytes: bytes) -> str:
        for encoding in ("utf-8", "latin-1", "cp1252"):
            try:
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="replace")

    @staticmethod
    def _parse_docx(file_bytes: bytes) -> str:
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(file_bytes))
        return "\n".join(para.text for para in doc.paragraphs if para.text.strip())

    @staticmethod
    def _parse_csv(file_bytes: bytes) -> str:
        import csv

        text_io = io.StringIO(file_bytes.decode("utf-8", errors="replace"))
        reader = csv.reader(text_io)
        rows = [", ".join(row) for row in reader]
        return "\n".join(rows)

    @staticmethod
    def _parse_xlsx(file_bytes: bytes) -> str:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                line = ", ".join(str(cell) if cell is not None else "" for cell in row)
                if line.strip(","):
                    lines.append(line)
        return "\n".join(lines)
